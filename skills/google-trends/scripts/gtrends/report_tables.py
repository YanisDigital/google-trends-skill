"""Tabular outputs: CSV files (stdlib) and an XLSX workbook (needs openpyxl)."""
from __future__ import annotations

import csv
from pathlib import Path

from .analyze import monthly
from .texts import VERDICT_TEXT
from .util import spreadsheet_safe as safe


def _write_csv(path: Path, header: list[str], rows: list[list]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        for r in rows:
            w.writerow([safe(c) for c in r])


def timeline_rows(snap: dict, source: str) -> tuple[list[str], list[list]]:
    tl = snap["sources"].get(source, {}).get("timeline")
    if not tl:
        return [], []
    labels = list(tl["series"])
    rows = [[d, *[tl["series"][lab][i] for lab in labels]] for i, d in enumerate(tl["dates"])]
    return ["date", *labels], rows


def region_rows(snap: dict) -> list[list]:
    rows = []
    for src, node in snap["sources"].items():
        for lab, regs in node.get("regions", {}).items():
            rows += [[src, lab, r["name"], r["value"]] for r in regs]
    return rows


def related_rows(snap: dict) -> list[list]:
    rows = []
    for src, node in snap["sources"].items():
        for lab, rel in node.get("related", {}).items():
            for kind in ("top", "rising"):
                for r in rel.get(kind, []):
                    rows.append([src, lab, kind, r["query"], "breakout" if r["breakout"] else r["value"]])
    return rows


def write_csvs(snap: dict, analysis: dict, folder: Path) -> list[Path]:
    written = []
    for src in snap["sources"]:
        header, rows = timeline_rows(snap, src)
        if rows:
            p = folder / f"timeline_{src}.csv"
            _write_csv(p, header, rows)
            written.append(p)
    rr = region_rows(snap)
    if rr:
        p = folder / "regions.csv"
        _write_csv(p, ["source", "term", "region", "value"], rr)
        written.append(p)
    rel = related_rows(snap)
    if rel:
        p = folder / "related.csv"
        _write_csv(p, ["source", "term", "kind", "query", "value"], rel)
        written.append(p)
    return written


def write_xlsx(snap: dict, analysis: dict, path: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:  # pragma: no cover - exercised only without the optional dependency
        raise RuntimeError("openpyxl is required for XLSX output: pip install -r requirements.txt") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Term", "Verdict", "Shape", "Level last 12m", "YoY", "Buying intent", "Peak month", "Next peak", "Ads start from", "Ads start until"])
    for lab, t in analysis["terms"].items():
        web = t["web"]
        seas, tm = web.get("seasonality") or {}, web.get("timing") or {}
        ws.append([safe(lab), VERDICT_TEXT[t["verdict"]["label"]], web["shape"], web.get("avg_last_12m"), web.get("yoy_ratio"),
                   t["intent"]["status"], seas.get("peak_month") if seas.get("is_seasonal") else None,
                   tm.get("next_peak"), tm.get("ads_start_from"), tm.get("ads_start_until")])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for src in snap["sources"]:
        header, rows = timeline_rows(snap, src)
        if rows:
            sheet = wb.create_sheet(f"Timeline {src}"[:31])
            sheet.append([safe(h) for h in header])
            for r in rows:
                sheet.append(r)
            monthly_sheet = wb.create_sheet(f"Monthly {src}"[:31])
            monthly_sheet.append(["month", *[safe(h) for h in header[1:]]])
            tl = snap["sources"][src]["timeline"]
            by_label = {lab: monthly(tl["dates"], tl["series"][lab]) for lab in tl["series"]}
            first = next(iter(by_label.values()))
            for i, (y, m, _) in enumerate(first):
                monthly_sheet.append([f"{y}-{m:02d}", *[round(by_label[lab][i][2], 1) for lab in by_label]])

    for name, header, rows in (("Regions", ["source", "term", "region", "value"], region_rows(snap)),
                               ("Related", ["source", "term", "kind", "query", "value"], related_rows(snap))):
        if rows:
            sheet = wb.create_sheet(name)
            sheet.append(header)
            for r in rows:
                sheet.append([safe(c) for c in r])

    first_term = analysis.get("primary")
    if first_term in analysis["terms"]:
        sheet = wb.create_sheet("Checklist")
        sheet.append(["Item", "Status"])
        for c in analysis["terms"][first_term]["checklist"]:
            sheet.append([c["item"], c["status"]])
    wb.save(path)
    return path
